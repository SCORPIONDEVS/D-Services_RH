#!/usr/bin/env python3
"""Exporta a planilha 'Empregados Ativos - D SERVICE.xlsm' para JSON em database/seed-data/.
Uso: python3 tools/export_xlsm.py caminho/planilha.xlsm
Requer: pip install openpyxl
"""
import sys, json, re, datetime as dt, warnings, pathlib
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.utils import column_index_from_string as ci

SRC = sys.argv[1] if len(sys.argv) > 1 else "Empregados_Ativos_-_D_SERVICE.xlsm"
OUT = pathlib.Path(__file__).resolve().parent.parent / "database" / "seed-data"
OUT.mkdir(parents=True, exist_ok=True)
wb = openpyxl.load_workbook(SRC, data_only=True)

def s(v):
    if v is None: return None
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.time): return v.strftime("%H:%M")
    if isinstance(v, float) and v.is_integer(): v = int(v)
    v = str(v).strip()
    return v if v and v not in ("-", "#N/A") and not v.startswith("#") else None

def d(v):
    if v is None: return None
    if isinstance(v, dt.datetime): return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.date): return v.strftime("%Y-%m-%d")
    t = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", t)
    if m: return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2}).*", t)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else None

def n(v):
    if v is None or v == "": return None
    try: return round(float(str(v).replace(",", ".").replace("%", "")), 2)
    except ValueError: return None

def digits(v):
    x = re.sub(r"\D", "", str(v)) if v is not None else ""
    return x or None

def save(name, data):
    (OUT / f"{name}.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{name}: {len(data) if hasattr(data,'__len__') else ''}")

# ---------------- Funcionários ----------------
ws = wb["CADASTRO - FUNCIONÁRIOS"]
C = lambda col, r: ws.cell(r, ci(col)).value
emps = []
for r in range(3, ws.max_row + 1):
    name = s(C("A", r))
    if not name: continue
    deps = []
    for i, cols in enumerate([("BG","BH","BI","BJ"),("BK","BL","BM","BN"),("BO","BP","BQ","BR"),("BS","BT","BU","BV")]):
        dn = s(C(cols[0], r))
        if dn: deps.append({"name": dn, "birth_date": d(C(cols[1], r)), "cpf": digits(C(cols[2], r)), "relationship": s(C(cols[3], r))})
    emps.append({
        "name": name, "company": s(C("C", r)), "team": s(C("D", r)), "esocial_code": s(C("E", r)),
        "admission_date": d(C("F", r)), "salary": n(C("I", r)), "category": s(C("J", r)),
        "job_title": s(C("K", r)), "cpf": digits(C("L", r)), "aso_admission_date": d(C("M", r)),
        "termination_date": d(C("N", r)), "termination_reason": s(C("O", r)), "aso_dismissal_date": d(C("P", r)),
        "cbo": s(C("Q", r)), "function_description": s(C("R", r)), "work_schedule": s(C("S", r)),
        "scale": s(C("T", r)), "leader": s(C("U", r)), "sector": s(C("V", r)), "pis": digits(C("W", r)),
        "rg": s(C("X", r)), "rg_uf": s(C("Y", r)), "rg_issuer": s(C("Z", r)), "rg_issue_date": d(C("AA", r)),
        "birth_date": d(C("AB", r)), "birth_city": s(C("AC", r)), "birth_uf": s(C("AD", r)), "birth_country": s(C("AE", r)),
        "ctps": s(C("AF", r)), "ctps_series": s(C("AG", r)), "ctps_uf": s(C("AH", r)), "ctps_issue_date": d(C("AI", r)),
        "sex": s(C("AJ", r)), "marital_status": s(C("AK", r)), "race": s(C("AL", r)), "education": s(C("AM", r)),
        "address": s(C("AN", r)), "address_number": s(C("AO", r)), "address_complement": s(C("AP", r)),
        "district": s(C("AQ", r)), "zip": digits(C("AR", r)), "city": s(C("AS", r)), "state": s(C("AT", r)),
        "phone_alt": s(C("AU", r)), "mobile": s(C("AV", r)), "email": s(C("AW", r)),
        "bank": s(C("AX", r)), "pix_key": s(C("AY", r)), "mother_name": s(C("AZ", r)), "father_name": s(C("BA", r)),
        "uniform_shirt": s(C("BB", r)), "uniform_pants": s(C("BC", r)), "uniform_shoes": s(C("BD", r)),
        "uniform_vest": s(C("BE", r)), "uniform_gloves": s(C("BF", r)),
        "vt_benefit": (s(C("BW", r)) or "").upper() == "SIM", "vt_company": s(C("BX", r)),
        "vt_trips_per_day": n(C("BY", r)), "vt_tariff": n(C("BZ", r)),
        "vr_benefit": (s(C("CC", r)) or "").upper() == "SIM", "vr_discount_value": n(C("CE", r)),
        "benefit_count_type": s(C("CG", r)),
        "vacation_start": d(C("CQ", r)), "vacation_end": d(C("CR", r)),
        "aso_periodic": [x for x in (d(C("CK", r)), d(C("CL", r))) if x],
        "dependents": deps,
    })
save("employees", emps)

# ---------------- Base: cargos / valores / empresas / parâmetros ----------------
wsb = wb["BASE - VALORES POR FUNÇÃO 2026"]
B = lambda a: wsb[a].value
positions, vr = {}, {}
for r in range(4, 14):
    if s(B(f"F{r}")): positions[s(B(f"F{r}"))] = {"name": s(B(f"F{r}")), "group": s(B(f"G{r}")), "salary": n(B(f"H{r}")), "vr_per_day": None}
for r in range(4, 13):
    if s(B(f"J{r}")):
        positions.setdefault(s(B(f"J{r}")), {"name": s(B(f"J{r}")), "group": None, "salary": None, "vr_per_day": None})["vr_per_day"] = n(B(f"K{r}"))
save("positions", list(positions.values()))
save("companies", [{"name": s(B(f"P{r}")), "cnpj": s(B(f"Q{r}"))} for r in range(4, 7) if s(B(f"P{r}"))])
save("settings", {
    "va_monthly_value": n(B("N4")), "assiduity_salary_cap": n(B("N7")), "vt_discount_percent": 6,
    "assiduity_rules": [{"label": s(B(f"M{r}")), "value": n(B(f"N{r}"))} for r in range(10, 14)],
    "staff_quota": [{"job_title": s(B(f"F{r}")), "quota": n(B(f"G{r}")), "note": s(B(f"K{r}"))} for r in range(21, 31) if s(B(f"F{r}"))],
})

# ---------------- Vagas ----------------
wsv = wb["VAGAS"]
vac = []
for r in range(3, wsv.max_row + 1):
    role = s(wsv.cell(r, 2).value)
    if not role: continue
    V = lambda c: wsv.cell(r, c).value
    vac.append({"number": r - 2, "job_title": role, "team": s(V(3)), "leader": s(V(4)), "reason": s(V(5)),
                "previous_employee": s(V(7)), "opened_at": d(V(8)), "stage": s(V(9)), "closed_at": d(V(10)),
                "salary": n(V(11)), "period": s(V(12)), "schedule": s(V(13)), "approved_candidate": s(V(14))})
save("job_openings", vac)

wsc = wb["VAGAS - CANDIDATOS"]
cands = []
for r in range(4, wsc.max_row + 1):
    nm = s(wsc.cell(r, 1).value)
    if not nm: continue
    V = lambda c: wsc.cell(r, c).value
    cands.append({"name": nm, "contact": s(V(2)), "job_number": int(V(3)) if str(V(3) or "").isdigit() else None,
                  "status": s(V(5)), "message_sent": (s(V(6)) or "").upper() == "SIM", "recruiter": s(V(7)),
                  "interview_date": d(V(8)), "interview_time": s(V(9)), "interview_accepted": s(V(10)),
                  "referral": s(V(11)), "source": s(V(12)), "notes": s(V(16)),
                  "invite_message": s(V(13)), "approved_message": s(V(14)), "rejected_message": s(V(15))})
save("candidates", cands)

# ---------------- Atestados / CID ----------------
wsa = wb["CADASTRO - ATESTADOS"]
cert = []
for r in range(2, wsa.max_row + 1):
    nm = s(wsa.cell(r, 1).value)
    if not nm: continue
    V = lambda c: wsa.cell(r, c).value
    cert.append({"employee_name": nm, "team": s(V(2)), "type": s(V(3)), "start_date": d(V(4)), "cid": s(V(5)),
                 "cid_description": s(V(6)), "days": int(n(V(7)) or 0) or None, "entry_time": s(V(9)), "exit_time": s(V(10)),
                 "form": s(V(12)), "worked_note": s(V(13)), "vt_discount": s(V(14)), "vr_discount": s(V(15)),
                 "va_discount": s(V(16)), "competence": d(V(17)), "status": s(V(18))})
save("certificates", cert)

wsci = wb["CIDs"]
cids = []
for r in range(4, wsci.max_row + 1):
    t = s(wsci.cell(r, 1).value)
    if not t: continue
    m = re.match(r"^([A-Z]\d{2}(?:\.\d+)?)\s+(.*)$", t)
    if m: cids.append({"code": m.group(1), "description": m.group(2)})
_seen=set(); cids=[c for c in cids if not (c["code"] in _seen or _seen.add(c["code"]))]
save("cids", cids)

# ---------------- Folha: ocorrências, coberturas, assiduidade ----------------
wso = wb["FOLHA - Ocorrência de Ponto"]
occ = []
for r in range(3, wso.max_row + 1):
    nm = s(wso.cell(r, 1).value)
    if not nm: continue
    V = lambda c: wso.cell(r, c).value
    occ.append({"employee_name": nm, "team": s(V(2)), "date": d(V(3)), "competence": d(V(4)), "kind": s(V(5)), "reason": s(V(6)), "status": s(V(7))})
save("time_occurrences", occ)

wsf = wb["FT COBERTURAS"]
cov = []
for r in range(3, wsf.max_row + 1):
    cn = s(wsf.cell(r, 3).value)
    if not cn: continue
    V = lambda c: wsf.cell(r, c).value
    cov.append({"absent_name": s(V(1)), "date": d(V(2)), "cover_name": cn, "amount": n(V(4)), "team": s(V(5)),
                "pay_date": d(V(6)), "status": s(V(7)), "finance_status": s(V(8))})
save("coverages", cov)

wsas = wb["Assiduidade EM ABERTO"]
months = {}
for c in range(11, 22):
    m = re.search(r"(\d{2})/(\d{4})", str(wsas.cell(2, c).value or ""))
    if m: months[c] = f"{m.group(2)}-{m.group(1)}-01"
att = []
for r in range(3, wsas.max_row + 1):
    nm = s(wsas.cell(r, 1).value)
    if not nm: continue
    for c, comp in months.items():
        v = n(wsas.cell(r, c).value)
        if v is not None:
            att.append({"employee_name": nm, "competence": comp, "amount": v})
save("attendance_bonuses", att)

# ---------------- Calendário RH ----------------
wsk = wb["CALENDÁRIO RH"]
cal = []
for r in range(2, 33):
    day = wsk.cell(r, 1).value
    if not isinstance(day, (int, float)): continue
    for col, area in ((2, "RH"), (3, "RH"), (4, "Financeiro")):
        t = s(wsk.cell(r, col).value)
        if t:
            for line in t.split("\n"):
                if line.strip(): cal.append({"day": int(day), "area": area, "title": line.strip()})
save("calendar_tasks", cal)

# ---------------- Armários ----------------
wsl = wb["CONTROLE DE ARMÁRIOS"]
lockers = []
for base, group in ((1, "MASCULINO (seg/limp)"), (6, "FEMININO (seg/limp)"), (11, "TERMINAL")):
    section = None
    for r in range(3, wsl.max_row + 1):
        a = wsl.cell(r, base).value
        if isinstance(a, str) and a.lower().startswith("armário"): section = a.strip(); continue
        if isinstance(a, (int, float)):
            lockers.append({"group": group, "section": section, "slot": int(a), "occupant": s(wsl.cell(r, base + 1).value), "team": s(wsl.cell(r, base + 2).value)})
save("lockers", lockers)

# ---------------- Escalas de turno ----------------
sched = []
for sheet, shift in (("Escala 1º TURNO 2026", 1), ("Escala 2º TURNO 2026", 2)):
    w = wb[sheet]
    heads = [str(w.cell(4, c).value or "").replace("\n", " ") for c in range(1, 17)]
    for r in range(5, w.max_row + 1):
        day = w.cell(r, 1).value
        if not isinstance(day, dt.datetime): continue
        for c in range(3, 15, 2):
            sched.append({"date": day.strftime("%Y-%m-%d"), "shift": shift, "slot": heads[c-1].replace("Colab. ", ""),
                          "person": s(w.cell(r, c).value), "leader_rotation": (s(w.cell(r, c+1).value) or "").upper() == "SIM"})
        for c in (15, 16):
            sched.append({"date": day.strftime("%Y-%m-%d"), "shift": shift, "slot": heads[c-1].replace(" LÍDER", ""),
                          "person": s(w.cell(r, c).value), "leader_rotation": c == 16})
save("schedule_entries", sched)


# ============ HISTÓRICO E DEMAIS ABAS ============
def cellval(v):
    if v is None: return None
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime("%Y-%m-%d")
    if isinstance(v, dt.time): return v.strftime("%H:%M")
    if isinstance(v, dt.timedelta):
        m = int(v.total_seconds() // 60); return f"{m // 60}:{m % 60:02d}"
    if isinstance(v, float) and v.is_integer(): return int(v)
    if isinstance(v, float): return round(v, 4)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v

MONTHS = {"JANEIRO":1,"FEVEREIRO":2,"MARÇO":3,"ABRIL":4,"MAIO":5,"JUNHO":6,"JULHO":7,"AGOSTO":8,"SETEMBRO":9,"OUTUBRO":10,"NOVEMBRO":11,"DEZEMBRO":12}

def hhmm(v):
    """Converte horas da planilha (timedelta, time, datetime 1899/1900 ou texto) em 'H:MM'."""
    if v is None or v == "": return None
    if isinstance(v, dt.timedelta):
        m = round(v.total_seconds() / 60); return f"{m // 60}:{m % 60:02d}"
    if isinstance(v, dt.datetime) and v.year <= 1900:
        base = dt.datetime(1899, 12, 30)
        m = round((v - base).total_seconds() / 60); return f"{m // 60}:{m % 60:02d}"
    if isinstance(v, dt.time): return v.strftime("%H:%M")
    return s(v)

# --- Benefícios VT/VR por colaborador × período ---
wsh = wb["BENEFÍCIOS - VT-VR"]
cur, col_block = None, {}
for c in range(8, wsh.max_column + 1):
    lab = s(wsh.cell(3, c).value)
    if lab: cur = lab
    col_block[c] = cur

def parse_block(label):
    m = re.match(r"([A-ZÇÃ]+)/(\d{4})(?: - (.*))?$", label or "")
    if not m or m.group(1) not in MONTHS: return None, None
    ym = f"{m.group(2)}-{MONTHS[m.group(1)]:02d}"
    suf = (m.group(3) or "").upper()
    if "1ª" in suf: return f"{ym}-1Q", ("corr" if "CORREÇÃO" in suf else "main")
    if "2ª" in suf: return f"{ym}-2Q", ("diff" if "DIFERENÇA" in suf else "main")
    return f"{ym}-M", "main"

def map_field(f, kind):
    u = f.upper()
    if u.startswith("DIAS-"): return "vt_days" if " VT" in u else "vr_days"
    if "COMPLEMENTO" in u: return "correction_vt" if u.startswith("VT") else "correction_vr"
    if u.endswith(" DIF"): return "correction_vr"
    if re.match(r"^VT-[12] ", u): return "vt_amount"
    if re.match(r"^VR-[12] ", u) and not u.endswith("CORRETO"): return "vr_amount"
    if u in ("MARÇO VT",): return "vt_amount"
    if u in ("MARÇO VR",): return "vr_amount"
    if u.startswith("MOTIVO"): return "correction_reason" if kind == "corr" else "reason"
    if u.startswith("AJUSTES (+)"): return "adjustment_plus"
    if u.startswith("AJUSTES (-)"): return "adjustment_minus"
    if u.startswith("AJUSTES (VT)"): return "adjustment_vt"
    if u.startswith("AJUSTES (VR)"): return "adjustment_vr"
    return None

NUMERIC = {"vt_days","vr_days","vt_amount","vr_amount","adjustment_plus","adjustment_minus","adjustment_vt","adjustment_vr","correction_vt","correction_vr"}
entries = {}
for r in range(5, wsh.max_row + 1):
    nm = s(wsh.cell(r, 1).value)
    if not nm: continue
    for c in range(14, wsh.max_column + 1):
        head = s(wsh.cell(4, c).value); v = cellval(wsh.cell(r, c).value)
        if not head or v is None or col_block.get(c) is None: continue
        period, kind = parse_block(col_block[c])
        if not period: continue
        e = entries.setdefault((nm, period), {"employee_name": nm, "period": period, "extra": {}})
        key = map_field(head, kind)
        if key in NUMERIC:
            e[key] = n(v) if not isinstance(v, str) else None
        elif key:
            e[key] = str(v)
        else:
            e["extra"][head] = v
out_entries = []
for e in entries.values():
    nums = [e.get(k) for k in NUMERIC]
    has = any(x for x in nums if x) or e.get("reason") or e.get("correction_reason") or any(x for x in e["extra"].values() if x not in (0, "0"))
    if has: out_entries.append(e)
save("benefit_entries", out_entries)

# --- Apontamento de folha (Segurança e Limpeza) ---
comp = None
payroll = []
for sheet, team, cols in (
    ("FOLHA - Apontamento Seg.", "SEGURANÇA", dict(vr="H", vt="I", va="J", bonus="K", func="L", total="M", cd="N", cdates="O", fdates="P", obs="Q")),
    ("FOLHA - Apontamento Limp.", "LIMPEZA", dict(vr="H", vt="I", va="J", bonus="K", func=None, total="L", cd="M", cdates="N", fdates="O", obs="P")),
):
    w = wb[sheet]
    month = str(w["R1"].value if team == "SEGURANÇA" else w["Q1"].value or "").zfill(2)
    year = 2026
    competence = f"{year}-{month}-01"
    for r in range(9, w.max_row + 1):
        nm = s(w.cell(r, 1).value)
        if not nm or nm.upper() == "TOTAL": continue
        g = lambda col: w[f"{col}{r}"].value if col else None
        payroll.append({
            "employee_name": nm, "team": team, "competence": competence, "job_title": s(w.cell(r, 2).value),
            "night_hours": hhmm(w.cell(r, 3).value), "absence_days": n(w.cell(r, 4).value),
            "partial_absence_hours": hhmm(w.cell(r, 5).value), "overtime_50": hhmm(w.cell(r, 6).value),
            "overtime_100": hhmm(w.cell(r, 7).value), "meal_voucher": n(g(cols["vr"])), "transport_voucher": n(g(cols["vt"])),
            "food_voucher": n(g(cols["va"])), "assiduity_bonus": n(g(cols["bonus"])), "function_accumulation": n(g(cols["func"])),
            "total_va": n(g(cols["total"])), "certificate_days": n(g(cols["cd"])), "certificate_dates": s(g(cols["cdates"])),
            "absence_dates": s(g(cols["fdates"])), "notes": s(g(cols["obs"])),
        })
save("payroll_entries", payroll)

# --- Liderança Terminal (oculta): dias de liderança e 10% por competência ---
wl = wb["FOLHA - Liderança Terminal"]
lead = []
pairs = []
for c in range(6, wl.max_column + 1):
    h = str(wl.cell(1, c).value or "")
    m = re.match(r"LID\.\s+(\d{2})/(\d{4})", h)
    if m: pairs.append((c, f"{m.group(2)}-{m.group(1)}-01"))
for r in range(2, wl.max_row + 1):
    nm = s(wl.cell(r, 1).value)
    if not nm: continue
    for c, comp_ in pairs:
        days, amt = n(wl.cell(r, c).value), n(wl.cell(r, c + 1).value)
        if days is None and amt is None: continue
        lead.append({"employee_name": nm, "job_title": s(wl.cell(r, 2).value), "salary": n(wl.cell(r, 3).value),
                     "slot": s(wl.cell(r, 4).value), "shift": s(wl.cell(r, 5).value), "competence": comp_,
                     "lead_days": days, "amount": amt})
save("leadership_bonuses", lead)

# --- Inconsistências de ponto (oculta) ---
wi = wb["FOLHA - Inconsistências"]
inc = []
for r in range(3, wi.max_row + 1):
    nm = s(wi.cell(r, 1).value)
    if not nm: continue
    V = lambda c: wi.cell(r, c).value
    inc.append({"employee_name": nm, "team": s(V(2)), "shift": s(V(3)), "date": d(V(4)), "competence": d(V(5)),
                "occurrence": s(V(6)), "kind": s(V(7)), "doubt": s(V(8)), "notes": s(V(9)), "answer": s(V(10)), "status": s(V(11))})
save("inconsistencies", inc)

# --- Padrão mensal de escala (12x36 / PAR-ÍMPAR) ---
wp = wb["ESCALA - Padrão de Turnos"]
mm = {"jan":1,"fev":2,"mar":3,"abr":4,"mai":5,"jun":6,"jul":7,"ago":8,"set":9,"out":10,"nov":11,"dez":12}
pcols = {}
for c in range(9, 19):
    m = re.match(r"([a-z]{3})/(\d{2})", str(wp.cell(3, c).value or ""))
    if m: pcols[c] = f"20{m.group(2)}-{mm[m.group(1)]:02d}"
pat = []
for r in range(4, wp.max_row + 1):
    nm = s(wp.cell(r, 6).value)
    if not nm: continue
    for c, ym in pcols.items():
        v = s(wp.cell(r, c).value)
        if v: pat.append({"employee_name": nm, "team": s(wp.cell(r, 7).value), "schedule": s(wp.cell(r, 8).value), "month": ym, "pattern": v})
save("shift_patterns", pat)

assign = []
for r in range(4, 23):
    lab = wp.cell(r, 25).value
    if not isinstance(lab, str): continue
    assign.append({"position": lab.strip(), "odd_person": s(wp.cell(r, 26).value), "even_person": s(wp.cell(r, 27).value)})
save("pattern_assignments", assign)

# --- Posições fixas do Terminal (Colab. A..L / folguistas) ---
slots = []
for sheet, shift in (("Escala 1º TURNO 2026", 1), ("Escala 2º TURNO 2026", 2)):
    w = wb[sheet]
    for r in range(5, w.max_row + 1):
        code, person, role = s(w.cell(r, 23).value), s(w.cell(r, 21).value), s(w.cell(r, 22).value)
        if code: slots.append({"shift": shift, "slot": code, "person": person, "job_title": role})
save("schedule_slots", slots)

# --- Jornadas por cargo (aba HORAS) ---
wsw = wb["HORAS 🧮"]
def tm(v):
    return v.strftime("%H:%M") if isinstance(v, (dt.datetime, dt.time)) else s(v)
shifts = []
for r in range(3, 13):
    nm = s(wsw.cell(r, 2).value)
    if not nm: continue
    shifts.append({"name": nm, "entry_time": tm(wsw.cell(r, 3).value), "exit_time": tm(wsw.cell(r, 4).value),
                   "hours_total": hhmm(wsw.cell(r, 5).value), "hours_worked": hhmm(wsw.cell(r, 6).value)})
save("work_shifts", shifts)
